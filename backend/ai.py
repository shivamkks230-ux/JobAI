"""AI services: resume parsing/scoring, resume improvement, interview coach (via Emergent LLM key)."""
import io
import os
import re
import json
import uuid
import logging

from emergentintegrations.llm.chat import LlmChat, UserMessage, TextDelta, StreamDone

logger = logging.getLogger("jobmatch.ai")
MODEL = ("openai", "gpt-5.4")
HONESTY = ("Never invent experience, qualifications, employers, dates, certifications or skills that are not "
           "present in the provided text. If information is missing, leave it empty. Respond with JSON only.")


async def llm_json(system: str, prompt: str) -> dict:
    chat = LlmChat(api_key=os.environ["EMERGENT_LLM_KEY"], session_id=f"jm-{uuid.uuid4().hex}",
                   system_message=system).with_model(*MODEL)
    out = []
    async for ev in chat.stream_message(UserMessage(text=prompt)):
        if isinstance(ev, TextDelta):
            out.append(ev.content)
        elif isinstance(ev, StreamDone):
            break
    text = "".join(out)
    m = re.search(r"\{.*\}", text, re.S)
    if not m:
        raise ValueError("AI returned an invalid response")
    return json.loads(m.group(0))


def extract_text(data: bytes, ext: str) -> str:
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages)[:20000]
        if ext == "docx":
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)[:20000]
        # legacy .doc: best-effort printable text extraction
        txt = data.decode("latin-1", errors="ignore")
        return " ".join(re.findall(r"[A-Za-z0-9@.,:;()+/#&\-' ]{4,}", txt))[:20000]
    except Exception as e:
        logger.warning("text extraction failed: %s", e)
        return ""


async def analyze_resume(text: str) -> dict:
    system = "You are an expert ATS resume parser and reviewer for the Indian job market. " + HONESTY
    prompt = f"""Parse and review this resume. Return JSON with exactly these keys:
{{"parsed": {{"name": "", "email": "", "phone": "", "headline": "", "education": [{{"degree": "", "institution": "", "year": ""}}],
"experience": [{{"title": "", "company": "", "duration": "", "summary": ""}}], "companies": [], "job_titles": [], "skills": [],
"certifications": [], "projects": [], "keywords": [], "industry": "", "years_of_experience": 0}},
"score": 0, "strengths": [], "missing_skills": [], "missing_keywords": [], "formatting_issues": [], "experience_gaps": [], "suggestions": []}}
score is 0-100 for overall resume quality & ATS readiness. missing_skills/keywords are commonly expected for the candidate's
target roles but absent. Keep each list item short (max 15 words).

RESUME TEXT:
{text[:12000]}"""
    return await llm_json(system, prompt)


IMPROVE_MODES = {
    "resume": "Improve the overall resume",
    "summary": "Rewrite the professional summary",
    "experience": "Improve experience bullet points (action verbs, impact, clarity)",
    "skills": "Improve organisation of the skills section",
    "ats": "Optimise for ATS (keywords, headings, formatting)",
}


async def improve_resume(mode: str, text: str, target_role: str | None) -> dict:
    system = "You are a professional resume editor. " + HONESTY + \
        " Only rephrase what exists; you may suggest the candidate ADD items only if they genuinely have them."
    prompt = f"""Task: {IMPROVE_MODES[mode]}. Target role: {target_role or 'not specified'}.
Return JSON: {{"summary": "one paragraph overview", "items": [{{"section": "", "before": "", "after": "", "why": ""}}], "tips": []}}
Give 3-6 items with exact before text from the resume and an improved after version.

RESUME TEXT:
{text[:10000]}"""
    return await llm_json(system, prompt)


async def interview_questions(category: str, role: str | None, level: str | None) -> dict:
    system = "You are an experienced interviewer in India. Respond with JSON only."
    prompt = f"""Create 5 realistic mock interview questions for category "{category}", role "{role or 'general'}",
experience level "{level or 'any'}". Mix behavioural and role-specific questions.
Return JSON: {{"questions": ["...", "..."]}}"""
    return await llm_json(system, prompt)


async def interview_feedback(category: str, question: str, answer: str) -> dict:
    system = ("You are a supportive interview coach. Evaluate only the text answer given. This is practice feedback, "
              "not a hiring decision. Respond with JSON only.")
    prompt = f"""Category: {category}
Question: {question}
Candidate answer: {answer[:4000]}
Return JSON: {{"scores": {{"confidence": 0, "fluency": 0, "clarity": 0, "relevance": 0, "vocabulary": 0, "structure": 0}},
"overall": 0, "strengths": [], "improvements": [], "sample_answer": ""}}
Scores are 0-10. Keep lists to 2-4 short items. sample_answer is a concise improved model answer (no invented facts)."""
    return await llm_json(system, prompt)
